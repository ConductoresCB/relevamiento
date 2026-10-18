"""
Genera la planilla base con el tablero automático.
Uso: python3 planilla/generar_planilla.py
  - planilla/Base_Relevamiento_Tablero.xlsx         (vacía, para usar)
  - planilla/Tablero_DEMO_datos_ficticios.xlsx      (con 126 relevamientos ficticios, para ver cómo queda)
Si cambian el cuestionario: node planilla/exportar_columnas.js 120 y volver a correr este script.
"""
import datetime, json, os
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
DEF = json.load(open(os.path.join(AQUI, 'columnas.json'), encoding='utf-8'))
COLS = DEF['columnas']
LETRA = {c['id']: get_column_letter(i + 1) for i, c in enumerate(COLS)}
ULTIMA = 3000  # filas de datos contempladas por las fórmulas

TINTA, AZUL, GRIS = '1C1E21', '1A588F', '4A4E54'
F = 'Arial'
f_base = Font(name=F, size=10)
f_titulo = Font(name=F, size=18, bold=True, color=TINTA)
f_sub = Font(name=F, size=10, italic=True, color=GRIS)
f_bloque = Font(name=F, size=12, bold=True, color=AZUL)
f_enc = Font(name=F, size=10, bold=True, color='FFFFFF')
f_neg = Font(name=F, size=10, bold=True)
f_input = Font(name=F, size=10, color='0000FF', bold=True)
relleno_enc = PatternFill('solid', fgColor=TINTA)
relleno_total = PatternFill('solid', fgColor='EDE9DF')
relleno_input = PatternFill('solid', fgColor='FFFF00')
relleno_kpi = PatternFill('solid', fgColor='E3EDF8')
linea = Side(style='thin', color='CFC8B8')
borde = Border(bottom=linea)
RAMPA_FR = ['EED59B', 'DEA553', 'CA7322', 'A44C1D', '772C1A']
RAMPA_FH = ['EED59B', 'CA7322', '772C1A']


def R(cid):
    l = LETRA[cid]
    return f'Datos!${l}$2:${l}${ULTIMA}'


def exacto_multiple(cid, cat):
    r = R(cid)
    return f'COUNTIF({r},"{cat}")+COUNTIF({r},"{cat};*")+COUNTIF({r},"*; {cat}")+COUNTIF({r},"*; {cat};*")'


def construir(demo_filas=None, aviso=None):
    """aviso: texto rojo bajo el título del tablero (si se pasa, reemplaza al de los datos ficticios)."""
    wb = Workbook()
    ins = wb.active; ins.title = 'Instrucciones'
    tab = wb.create_sheet('Tablero')
    dat = wb.create_sheet('Datos')
    con = wb.create_sheet('Contactos')
    dic = wb.create_sheet('Diccionario')

    # ---------------- Datos ----------------
    for i, c in enumerate(COLS, start=1):
        celda = dat.cell(row=1, column=i, value=c['id'])
        celda.font = f_enc; celda.fill = relleno_enc
        celda.comment = Comment(c['texto'], 'Instrumento')
        dat.column_dimensions[get_column_letter(i)].width = max(12, min(28, len(c['id']) + 4))
    dat.freeze_panes = 'D2'
    dat.sheet_properties.tabColor = AZUL
    if demo_filas:
        for fila_n, f in enumerate(demo_filas, start=2):
            for i, c in enumerate(COLS, start=1):
                v = f.get(c['id'], '')
                if v == '':
                    continue
                if c['id'] == 'fecha':
                    v = datetime.date.fromisoformat(v)
                    dat.cell(row=fila_n, column=i, value=v).number_format = 'dd/mm/yyyy'
                    continue
                dat.cell(row=fila_n, column=i, value=v).font = f_base

    # ---------------- Contactos ----------------
    for i, cid in enumerate(DEF['contacto'], start=1):
        celda = con.cell(row=1, column=i, value=cid); celda.font = f_enc; celda.fill = relleno_enc
        con.column_dimensions[get_column_letter(i)].width = 18
    con['A1'].comment = Comment('Datos personales. No compartir esta hoja: protegerla (Datos > Proteger hojas) '
                                'y no incluirla en exportaciones para publicar.', 'Instrumento')
    con.freeze_panes = 'A2'
    con.sheet_properties.tabColor = '9A3412'

    # ---------------- Diccionario ----------------
    enc = ['Variable', 'Descripción o pregunta', 'Sección', 'Valores posibles', 'Ejemplo']
    for i, t in enumerate(enc, start=1):
        c = dic.cell(row=1, column=i, value=t); c.font = f_enc; c.fill = relleno_enc
    ej = DEF['ejemplo']
    for n, c in enumerate(COLS, start=2):
        valores = '; '.join(c['opciones']) if c.get('opciones') else ('Texto libre' if c.get('tipo') == 'texto' else '')
        if c['id'] in ('fr_categoria',):
            valores = 'Bajo; Ligeramente elevado; Moderado; Alto; Muy alto'
        if c['id'] in ('fh_categoria',):
            valores = 'Bajo; Moderado; Alto'
        fila = [c['id'], c['texto'], c['seccion'], valores, ej.get(c['id'], '')]
        for i, v in enumerate(fila, start=1):
            celda = dic.cell(row=n, column=i, value=v); celda.font = f_base
            celda.alignment = Alignment(vertical='top', wrap_text=i in (2, 4))
    for col, w in zip('ABCDE', (24, 60, 22, 50, 22)):
        dic.column_dimensions[col].width = w
    dic.freeze_panes = 'B2'
    nota = dic.cell(row=len(COLS) + 3, column=1, value='La columna «Ejemplo» muestra el relevamiento de prueba 010-44-4600 (valores ficticios).')
    nota.font = f_sub

    # ---------------- Tablero ----------------
    tab.sheet_view.showGridLines = False
    for col, w in zip('ABCDEFG', (2, 42, 15, 14, 14, 14, 3)):
        tab.column_dimensions[col].width = w
    tab['B1'] = 'Tablero · Conductores de camiones del Corredor Bioceánico (Jujuy)'; tab['B1'].font = f_titulo
    if aviso:
        tab['B2'] = aviso
        tab['B2'].font = Font(name=F, size=10, bold=True, color='9A3412')
    elif demo_filas:
        tab['B2'] = 'DATOS FICTICIOS DE EJEMPLO. Para usar la planilla real, partí de Base_Relevamiento_Tablero.xlsx.'
        tab['B2'].font = Font(name=F, size=10, bold=True, color='9A3412')
    else:
        tab['B2'] = 'Se actualiza solo con lo que llega a la hoja «Datos». No escribir en esta hoja salvo las celdas amarillas.'
        tab['B2'].font = f_sub
    fila = 4

    def encabezado(fila, textos):
        for i, t in enumerate(textos):
            c = tab.cell(row=fila, column=2 + i, value=t); c.font = f_enc; c.fill = relleno_enc
            c.alignment = Alignment(horizontal='left' if i == 0 else 'center', wrap_text=True)

    def grafico(titulo, fila_ini, fila_fin, col_datos=3, colores=None, pct=False):
        ch = BarChart(); ch.type = 'bar'; ch.style = 10; ch.title = titulo
        datos = Reference(tab, min_col=col_datos, min_row=fila_ini, max_row=fila_fin)
        cats = Reference(tab, min_col=2, min_row=fila_ini, max_row=fila_fin)
        ch.add_data(datos, titles_from_data=False); ch.set_categories(cats)
        ch.legend = None
        ch.x_axis.scaling.orientation = 'maxMin'
        ch.x_axis.delete = False; ch.y_axis.delete = False
        ch.y_axis.majorGridlines = None
        if pct:
            ch.y_axis.numFmt = '0%'; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1
        s = ch.series[0]
        s.graphicalProperties.solidFill = AZUL; s.graphicalProperties.line.solidFill = AZUL
        if colores:
            for i, col in enumerate(colores):
                p = DataPoint(idx=i); p.graphicalProperties.solidFill = col; p.graphicalProperties.line.solidFill = col
                s.dPt.append(p)
        s.dLbls = DataLabelList(); s.dLbls.showVal = True
        for a in ('showSerName', 'showCatName', 'showLegendKey', 'showPercent'):
            setattr(s.dLbls, a, False)
        if pct:
            s.dLbls.numFmt = '0%'
        n = fila_fin - fila_ini + 1
        ch.width = 16; ch.height = max(6.0, 1.6 + 0.8 * n)
        return ch, int(ch.height / 0.53) + 1

    def bloque(titulo, cid, categorias, multiple=False, nota=None, colores=None, por_sexo=False, chart=True):
        nonlocal fila
        t = tab.cell(row=fila, column=2, value=titulo); t.font = f_bloque
        inicio_bloque = fila
        fila += 1
        if nota:
            n_ = tab.cell(row=fila, column=2, value=nota); n_.font = f_sub; fila += 1
        cab = ['Categoría', 'n', '%'] + (['Varón (n)', 'Mujer (n)'] if por_sexo else [])
        encabezado(fila, cab); fila += 1
        f_ini = fila
        f_tot = fila + len(categorias)
        for cat in categorias:
            tab.cell(row=fila, column=2, value=cat).font = f_base
            cnt = exacto_multiple(cid, cat) if multiple else f'COUNTIF({R(cid)},"{cat}")'
            tab.cell(row=fila, column=3, value='=' + cnt).font = f_base
            c = tab.cell(row=fila, column=4, value=f'=IF($C${f_tot}=0,"",C{fila}/$C${f_tot})'); c.number_format = '0.0%'; c.font = f_base
            if por_sexo:
                for k, sx in enumerate(('Varón', 'Mujer')):
                    tab.cell(row=fila, column=5 + k, value=f'=COUNTIFS({R(cid)},"{cat}",{R("b_sexo")},"{sx}")').font = f_base
            for col in range(2, 5 + (2 if por_sexo else 0)):
                tab.cell(row=fila, column=col).border = borde
            fila += 1
        tab.cell(row=fila, column=2, value='Respondieron' if not multiple else 'Respondieron (puede elegirse más de una)').font = f_neg
        tab.cell(row=fila, column=3, value=f'=COUNTA({R(cid)})').font = f_neg
        for col in range(2, 5 + (2 if por_sexo else 0)):
            tab.cell(row=fila, column=col).fill = relleno_total
        if por_sexo:
            for k in range(2):
                l = get_column_letter(5 + k)
                tab.cell(row=fila, column=5 + k, value=f'=SUM({l}{f_ini}:{l}{f_tot - 1})').font = f_neg
        fin_tabla = fila
        filas_graf = 0
        if chart:
            ch, filas_graf = grafico(titulo, f_ini, f_tot - 1, colores=colores)
            tab.add_chart(ch, f'I{inicio_bloque}')
        fila = max(fin_tabla, inicio_bloque + filas_graf) + 2

    # Indicadores generales
    tab.cell(row=fila, column=2, value='Indicadores generales').font = f_bloque; fila += 1
    encabezado(fila, ['Indicador', 'Valor']); fila += 1
    kpis = [
        ('Relevamientos recibidos', f'=COUNTA({R("id")})', '0'),
        ('Participantes únicos (por ID)', f'=IFERROR(SUMPRODUCT(({R("participante_id")}<>"")/COUNTIF({R("participante_id")},{R("participante_id")}&"")),0)', '0'),
        ('Relevamientos finalizados', f'=COUNTIF({R("estado")},"finalizado")', '0'),
        ('Con sociodemográfica completa', f'=COUNTIF({R("socio_completo")},"Sí")', '0'),
        ('Con biológica completa', f'=COUNTIF({R("bio_completo")},"Sí")', '0'),
        ('Con glucometría realizada', f'=COUNTIF({R("gluc_realizada")},"Sí")', '0'),
        ('Primer relevamiento', f'=IF(COUNT({R("fecha")})=0,"",MIN({R("fecha")}))', 'dd/mm/yyyy'),
        ('Último relevamiento', f'=IF(COUNT({R("fecha")})=0,"",MAX({R("fecha")}))', 'dd/mm/yyyy'),
    ]
    for etiqueta, formula, formato in kpis:
        tab.cell(row=fila, column=2, value=etiqueta).font = f_base
        c = tab.cell(row=fila, column=3, value=formula); c.number_format = formato; c.font = f_neg; c.fill = relleno_kpi
        tab.cell(row=fila, column=2).border = borde
        fila += 1
    fila += 1

    # Avance por equipo (nombres editables)
    tab.cell(row=fila, column=2, value='Avance por equipo').font = f_bloque; fila += 1
    tab.cell(row=fila, column=2, value='Las celdas amarillas son editables: escribí la letra de cada equipo tal como se cargó en Ajustes.').font = f_sub; fila += 1
    encabezado(fila, ['Equipo', 'Relevamientos', 'Finalizados', 'Con biológica', 'Último']); fila += 1
    for eq in ['A', 'B', 'C', 'D', 'E']:
        c = tab.cell(row=fila, column=2, value=eq); c.font = f_input; c.fill = relleno_input
        tab.cell(row=fila, column=3, value=f'=IF(B{fila}="","",COUNTIF({R("equipo")},B{fila}))').font = f_base
        tab.cell(row=fila, column=4, value=f'=IF(B{fila}="","",COUNTIFS({R("equipo")},B{fila},{R("estado")},"finalizado"))').font = f_base
        tab.cell(row=fila, column=5, value=f'=IF(B{fila}="","",COUNTIFS({R("equipo")},B{fila},{R("bio_completo")},"Sí"))').font = f_base
        c = tab.cell(row=fila, column=6, value=f'=IF(OR(B{fila}="",COUNTIF({R("equipo")},B{fila})=0),"",_xlfn.MAXIFS({R("fecha")},{R("equipo")},B{fila}))')
        c.number_format = 'dd/mm/yyyy'; c.font = f_base
        fila += 1
    fila += 1

    # Promedios
    tab.cell(row=fila, column=2, value='Promedios').font = f_bloque; fila += 1
    tab.cell(row=fila, column=2, value='La glucemia promedia mediciones con distintos tiempos de ayuno: usarla solo como referencia.').font = f_sub; fila += 1
    encabezado(fila, ['Variable', 'Total', 'Varón', 'Mujer', 'n']); fila += 1
    for etiqueta, cid in [('Edad (años)', 'edad'), ('IMC (kg/m²)', 'b_imc'), ('Perímetro abdominal (cm)', 'b_perimetro_cm'),
                          ('Presión sistólica (mmHg)', 'b_pas_mmhg'), ('Glucemia periférica (mg/dL)', 'g_resultado_mgdl'),
                          ('Años como chofer', 'p10_anios_chofer'), ('Horas de conducción sin paradas', 'p16_horas_sin_parar'),
                          ('Horas de sueño por día', 'p19_horas_sueno')]:
        tab.cell(row=fila, column=2, value=etiqueta).font = f_base
        c = tab.cell(row=fila, column=3, value=f'=IFERROR(AVERAGE({R(cid)}),"")'); c.number_format = '0.0'; c.font = f_base
        for k, sx in enumerate(('Varón', 'Mujer')):
            c = tab.cell(row=fila, column=4 + k, value=f'=IFERROR(AVERAGEIFS({R(cid)},{R("b_sexo")},"{sx}"),"")'); c.number_format = '0.0'; c.font = f_base
        tab.cell(row=fila, column=6, value=f'=COUNT({R(cid)})').font = f_base
        for col in range(2, 7):
            tab.cell(row=fila, column=col).border = borde
        fila += 1
    fila += 1

    bloque('Riesgo de diabetes a 10 años (FINDRISK)', 'fr_categoria', ['Bajo', 'Ligeramente elevado', 'Moderado', 'Alto', 'Muy alto'],
           nota='Solo personas sin diagnóstico de diabetes y con sexo mujer o varón.', colores=RAMPA_FR, por_sexo=True)
    bloque('Riesgo cardiovascular a 10 años (Framingham con IMC)', 'fh_categoria', ['Bajo', 'Moderado', 'Alto'],
           nota='Bajo: menos de 10 %. Moderado: 10 a 19 %. Alto: 20 % o más. Edades de 20 a 79 años.', colores=RAMPA_FH, por_sexo=True)

    # IMC por categorías (OMS)
    t = tab.cell(row=fila, column=2, value='Estado nutricional según IMC'); t.font = f_bloque; ini_imc = fila; fila += 1
    tab.cell(row=fila, column=2, value='Cortes de la OMS para adultos: <18,5 bajo peso; 18,5–24,9 normal; 25–29,9 sobrepeso; ≥30 obesidad.').font = f_sub; fila += 1
    encabezado(fila, ['Categoría', 'n', '%']); fila += 1
    r = R('b_imc'); f_i = fila; f_t = fila + 4
    for cat, crit in [('Bajo peso', f'COUNTIFS({r},"<18.5")'), ('Normal', f'COUNTIFS({r},">=18.5",{r},"<25")'),
                      ('Sobrepeso', f'COUNTIFS({r},">=25",{r},"<30")'), ('Obesidad', f'COUNTIFS({r},">=30")')]:
        tab.cell(row=fila, column=2, value=cat).font = f_base
        tab.cell(row=fila, column=3, value='=' + crit).font = f_base
        c = tab.cell(row=fila, column=4, value=f'=IF($C${f_t}=0,"",C{fila}/$C${f_t})'); c.number_format = '0.0%'
        for col in range(2, 5):
            tab.cell(row=fila, column=col).border = borde
        fila += 1
    tab.cell(row=fila, column=2, value='Con IMC registrado').font = f_neg
    tab.cell(row=fila, column=3, value=f'=COUNT({r})').font = f_neg
    for col in range(2, 5):
        tab.cell(row=fila, column=col).fill = relleno_total
    ch, nf = grafico('Estado nutricional según IMC', f_i, f_t - 1, colores=['EED59B', '1A588F', 'CA7322', '772C1A'])
    tab.add_chart(ch, f'I{ini_imc}')
    fila = max(fila, ini_imc + nf) + 2

    # Prevalencias (respuestas Sí)
    t = tab.cell(row=fila, column=2, value='Prevalencias (respuesta «Sí»)'); t.font = f_bloque; ini_p = fila; fila += 1
    encabezado(fila, ['Indicador', 'Sí (n)', 'Respondieron', '%']); fila += 1
    f_i = fila
    for etiqueta, cid in [('Fuma tabaco', 'b_fuma'), ('Coquea', 'p36_coquea'), ('Consume alcohol (fuera del trabajo)', 'p37_alcohol'),
                          ('Diagnóstico de hipertensión', 'b_hta_dx'), ('Diagnóstico de diabetes', 'b_dbt_dx'),
                          ('Glucosa elevada alguna vez', 'b_glucosa_previa'), ('Dolores corporales frecuentes', 'p30_dolores'),
                          ('Siniestro vial en el último año', 'p24_siniestro'), ('Cree que su trabajo es un riesgo para su salud', 'p15_riesgo_salud'),
                          ('Actividad física diaria (≥30 min)', 'b_act_fisica'), ('Frutas o verduras todos los días', 'b_fyv_diario')]:
        tab.cell(row=fila, column=2, value=etiqueta).font = f_base
        tab.cell(row=fila, column=3, value=f'=COUNTIF({R(cid)},"Sí")').font = f_base
        tab.cell(row=fila, column=4, value=f'=COUNTA({R(cid)})').font = f_base
        c = tab.cell(row=fila, column=5, value=f'=IF(D{fila}=0,"",C{fila}/D{fila})'); c.number_format = '0.0%'; c.font = f_base
        for col in range(2, 6):
            tab.cell(row=fila, column=col).border = borde
        fila += 1
    ch, nf = grafico('Prevalencias (% de quienes respondieron)', f_i, fila - 1, col_datos=5, pct=True)
    tab.add_chart(ch, f'I{ini_p}')
    fila = max(fila, ini_p + nf) + 2

    Q = {c['id']: c for c in COLS}
    bloque('Lugar de nacimiento', 'p01_nac_pais', Q['p01_nac_pais']['opciones'])
    bloque('Máximo nivel educativo alcanzado', 'p02_educacion', Q['p02_educacion']['opciones'])
    bloque('Recorrido que realiza', 'p11_recorrido', Q['p11_recorrido']['opciones'])
    bloque('Relación laboral', 'p13_relacion_laboral', Q['p13_relacion_laboral']['opciones'])
    bloque('Estado de las rutas en Jujuy', 'p09_rutas_jujuy', Q['p09_rutas_jujuy']['opciones'])
    bloque('Con lo que gana, ¿llega a fin de mes?', 'p25_fin_de_mes', Q['p25_fin_de_mes']['opciones'])
    bloque('Cobertura médica', 'p26_cobertura', Q['p26_cobertura']['opciones'], multiple=True,
           nota='Se puede marcar más de una opción: los porcentajes suman más de 100 %.')
    bloque('Dónde come durante los viajes', 'p18_donde_come', Q['p18_donde_come']['opciones'], multiple=True,
           nota='Se puede marcar más de una opción: los porcentajes suman más de 100 %.')
    bloque('Bebidas energéticas en viaje', 'p35_energeticas', Q['p35_energeticas']['opciones'], multiple=True,
           nota='Se puede marcar más de una opción: los porcentajes suman más de 100 %.')

    for fila_t in range(1, fila + 1):
        for col in range(2, 8):
            c = tab.cell(row=fila_t, column=col)
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, size=c.font.size if c.font else 10, bold=c.font.bold if c.font else False,
                              color=c.font.color if c.font else None, italic=c.font.italic if c.font else False)
    tab.freeze_panes = 'A3'

    # ---------------- Instrucciones ----------------
    ins.sheet_view.showGridLines = False
    ins.column_dimensions['A'].width = 2; ins.column_dimensions['B'].width = 120
    textos = [
        ('Base de relevamiento · Corredor Bioceánico (Jujuy)', f_titulo),
        ('Planilla que recibe los datos de la app de campo y arma el tablero sola.', f_sub),
        ('', None),
        ('Hojas', f_bloque),
        ('Tablero: indicadores, tablas y gráficos. Se recalcula solo; las únicas celdas para escribir son las amarillas (nombres de equipos).', f_base),
        ('Datos: una fila por relevamiento. La llena el script o se pega el CSV de la app. No cambiar la fila 1 ni insertar columnas en el medio.', f_base),
        ('Contactos: correo y celular. Datos personales: proteger la hoja y no compartirla.', f_base),
        ('Diccionario: qué significa cada columna, sus valores posibles y un ejemplo.', f_base),
        ('', None),
        ('Opción 1 — Planilla de Google (recomendada, sin costo)', f_bloque),
        ('1. Subí este archivo al Drive de la cuenta de Gmail del proyecto y abrilo con Hojas de cálculo de Google (Archivo > Guardar como Hojas de cálculo de Google).', f_base),
        ('2. En esa planilla: Extensiones > Apps Script. Borrá lo que aparece, pegá el contenido de Codigo.gs y cambiá la clave de la línea TOKEN.', f_base),
        ('3. Implementar > Nueva implementación > tipo «Aplicación web». Ejecutar como: Yo. Quién tiene acceso: Cualquier usuario. Autorizá los permisos.', f_base),
        ('4. Copiá la URL que termina en /exec y pegala, junto con la clave, en Ajustes de cada celular. Tocá «Probar conexión».', f_base),
        ('5. Compartí la planilla con el equipo como Lector para que vean el tablero. Para descargarla en Excel: Archivo > Descargar > Microsoft Excel.', f_base),
        ('', None),
        ('Opción 2 — Solo Excel, sin internet', f_bloque),
        ('1. En la computadora de coordinación, la app importa los paquetes de todos los celulares (Ajustes > Importar paquetes) y descarga un CSV consolidado.', f_base),
        ('2. En Excel: Datos > Obtener datos > Desde texto/CSV, origen UTF-8 y delimitador punto y coma. Copiá las filas (sin el encabezado) y pegalas desde la celda A2 de la hoja Datos.', f_base),
        ('3. El tablero se actualiza solo. Para un nuevo consolidado, seleccioná las filas de Datos desde la 2, apretá Supr (no elimines filas) y pegá el CSV completo otra vez (ya viene sin duplicados).', f_base),
        ('', None),
        ('Para vaciar la planilla (por ejemplo después de una prueba): seleccioná desde la fila 2 de Datos y de Contactos y apretá Supr. Nunca borres la fila 1.', f_base),
        ('', None),
        ('Notas metodológicas', f_bloque),
        (f'Las fórmulas cubren hasta la fila {ULTIMA} de Datos. Si se superan, ampliar los rangos.', f_base),
        ('FINDRISK y Framingham se calculan en la app con las tablas del protocolo; la planilla solo cuenta categorías.', f_base),
        ('IMC exactamente 30: FINDRISK lo puntúa como «25 a 30» (1 punto) y Framingham como «30 o más», tal como dice el modelo de encuesta. Revisar el criterio.', f_base),
        ('Las preguntas de opción múltiple se guardan separadas por «; » (por ejemplo: «Español; Portugués»).', f_base),
    ]
    for i, (t, fnt) in enumerate(textos, start=1):
        c = ins.cell(row=i, column=2, value=t)
        if fnt:
            c.font = fnt
        c.alignment = Alignment(wrap_text=True, vertical='top')
    ins.sheet_properties.tabColor = 'E8A317'
    wb.active = 1 if demo_filas else 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and (c.font is None or c.font.name != F):
                    c.font = Font(name=F, size=10)
    return wb


if __name__ == '__main__':
    base = os.path.join(AQUI, 'Base_Relevamiento_Tablero.xlsx')
    construir().save(base); print('→', base)
    demo = json.load(open(os.path.join(AQUI, 'demo_filas.json'), encoding='utf-8'))
    ruta_demo = os.path.join(AQUI, 'Tablero_DEMO_datos_ficticios.xlsx')
    construir(demo).save(ruta_demo); print('→', ruta_demo, f'({len(demo)} filas ficticias)')
